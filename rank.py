import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---- paths ----
INPUT = Path(r"C:\Users\Lemuel Quaye\Documents\School\students_info\whatiwant.json")
OUTPUT = INPUT.with_name("students_ranking.xlsx")

# ---- load ----
with open(INPUT, "r", encoding="utf-8") as f:
    raw = json.load(f)

students = []
courses = {}  # code -> (name, credit)

for student_id, rec in raw.items():
    if not rec.get("isSuccessful") or not rec.get("data"):
        continue
    sem = rec["data"][0]
    sheets = {}
    for s in sem["sheets"]:
        code = s["code"]
        courses[code] = (s["courseName"].strip(), s["credit"])
        sheets[code] = s
    students.append({
        "id": student_id,
        "name": rec["studentName"].strip(),
        "cwa": sem["cwa"],
        "reg": sem["creditRegistered"],
        "earned": sem["creditEarned"],
        "fails": sum(1 for s in sem["sheets"] if not s["hasPassed"]),
        "sheets": sheets,
    })

# ---- rank (highest CWA first; ties share a rank) ----
students.sort(key=lambda s: (-s["cwa"], s["name"]))
prev, rank = None, 0
for i, s in enumerate(students, start=1):
    if s["cwa"] != prev:
        rank = i
        prev = s["cwa"]
    s["rank"] = rank

course_codes = sorted(courses)

# ---- styling helpers ----
FONT = "Arial"
head_font = Font(name=FONT, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="1F4E78")
body_font = Font(name=FONT)
bold_font = Font(name=FONT, bold=True)
fail_fill = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header(ws, row=1):
    for c in ws[row]:
        c.font, c.fill, c.alignment, c.border = head_font, head_fill, center, border


def style_body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for c in row:
            c.font, c.border = body_font, border
            if c.column > 1:
                c.alignment = Alignment(horizontal="center")


def autosize(ws, min_w=8, max_w=40):
    for col in ws.columns:
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_w, min(longest + 2, max_w))


wb = Workbook()

# ---- Sheet 1: Ranking (total score per course) ----
ws = wb.active
ws.title = "Ranking"
ws.append(
    ["Rank", "Student ID", "Name", "CWA", "Credits Registered", "Credits Earned", "Failed Courses"]
    + [f"{code}\n{courses[code][0]}\n({courses[code][1]} cr)" for code in course_codes]
)
for s in students:
    ws.append(
        [s["rank"], s["id"], s["name"], s["cwa"], s["reg"], s["earned"], s["fails"]]
        + [s["sheets"][c]["fullScore"] if c in s["sheets"] else None for c in course_codes]
    )
style_header(ws)
style_body(ws)
ws.row_dimensions[1].height = 60
autosize(ws)
ws.column_dimensions["C"].width = 40
for r in range(2, ws.max_row + 1):
    ws.cell(r, 3).alignment = Alignment(horizontal="left")
    ws.cell(r, 4).number_format = "0.00"
    for ci, code in enumerate(course_codes, start=8):
        cell = ws.cell(r, ci)
        cell.number_format = "0.00"
        sheet = students[r - 2]["sheets"].get(code)
        if sheet and not sheet["hasPassed"]:
            cell.fill = fail_fill
ws.freeze_panes = "D2"
ws.auto_filter.ref = ws.dimensions

# ---- Sheet 2: Detailed scores (class / exam / total / grade per course) ----
wd = wb.create_sheet("Detailed Scores")
header = ["Rank", "Student ID", "Name", "CWA"]
for code in course_codes:
    header += [f"{code} Class", f"{code} Exam", f"{code} Total", f"{code} Grade"]
wd.append(header)
for s in students:
    row = [s["rank"], s["id"], s["name"], s["cwa"]]
    for code in course_codes:
        sh = s["sheets"].get(code)
        row += (
            [sh["classScore"], sh["examScore"], sh["fullScore"], sh["letter"]]
            if sh else [None] * 4
        )
    wd.append(row)
style_header(wd)
style_body(wd)
wd.row_dimensions[1].height = 32
autosize(wd, max_w=14)
wd.column_dimensions["C"].width = 40
for r in range(2, wd.max_row + 1):
    wd.cell(r, 3).alignment = Alignment(horizontal="left")
    for c in range(4, wd.max_column + 1):
        if isinstance(wd.cell(r, c).value, float):
            wd.cell(r, c).number_format = "0.00"
    for i, code in enumerate(course_codes):
        sh = students[r - 2]["sheets"].get(code)
        if sh and not sh["hasPassed"]:
            for off in range(4):
                wd.cell(r, 5 + i * 4 + off).fill = fail_fill
wd.freeze_panes = "E2"
wd.auto_filter.ref = wd.dimensions

# ---- Sheet 3: Long format (one row per student per course) ----
wl = wb.create_sheet("All Records")
wl.append(["Rank", "Student ID", "Name", "Course Code", "Course Name", "Credit",
           "Class Score", "Exam Score", "Total", "Grade", "Description", "Passed"])
for s in students:
    for code in course_codes:
        sh = s["sheets"].get(code)
        if not sh:
            continue
        wl.append([s["rank"], s["id"], s["name"], code, sh["courseName"].strip(), sh["credit"],
                   sh["classScore"], sh["examScore"], sh["fullScore"], sh["letter"],
                   sh["descriptions"], "Yes" if sh["hasPassed"] else "No"])
style_header(wl)
style_body(wl)
autosize(wl)
for r in range(2, wl.max_row + 1):
    wl.cell(r, 3).alignment = Alignment(horizontal="left")
    wl.cell(r, 5).alignment = Alignment(horizontal="left")
    if wl.cell(r, 12).value == "No":
        for c in range(1, 13):
            wl.cell(r, c).fill = fail_fill
wl.freeze_panes = "D2"
wl.auto_filter.ref = wl.dimensions

wb.save(OUTPUT)
print(f"Done. {len(students)} students ranked -> {OUTPUT}")